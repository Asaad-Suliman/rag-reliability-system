"""DOCX parsing via python-docx.

Real pagination is a Word rendering-time concept, not stored data — except
manual page breaks (Ctrl+Enter / `Document.add_page_break()`), which *are*
stored as `<w:br w:type="page"/>` runs. Where present we split on those;
otherwise pages are synthetic fixed-size windows, same as TXT/MD.
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.services.parsing.base import (
    ParsedDocument,
    ParsedPage,
    ParseError,
    ParseErrorKind,
    synthetic_paginate,
)


def _has_page_break(paragraph: Paragraph) -> bool:
    for run in paragraph.runs:
        for br in run._element.findall(qn("w:br")):
            if br.get(qn("w:type")) == "page":
                return True
    return False


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = docx.Document(str(path))
            paragraphs = [p.text for p in document.paragraphs]
            break_after = {i for i, p in enumerate(document.paragraphs) if _has_page_break(p)}
        except Exception as exc:
            raise ParseError(
                ParseErrorKind.CORRUPT, "The file could not be read; it may be corrupt."
            ) from exc

        if not break_after:
            full_text = "\n\n".join(paragraphs)
            return ParsedDocument(pages=synthetic_paginate(full_text), pages_are_synthetic=True)

        pages: list[ParsedPage] = []
        current: list[str] = []
        page_number = 1
        for i, text in enumerate(paragraphs):
            current.append(text)
            if i in break_after:
                pages.append(ParsedPage(page_number=page_number, text="\n\n".join(current)))
                page_number += 1
                current = []
        if current:
            pages.append(ParsedPage(page_number=page_number, text="\n\n".join(current)))
        return ParsedDocument(pages=pages, pages_are_synthetic=False)


def _demo() -> None:
    import tempfile

    parser = DocxParser()

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        # Explicit page break → two real pages.
        with_break = root / "with_break.docx"
        d = docx.Document()
        d.add_paragraph("page one content")
        d.add_page_break()  # type: ignore[no-untyped-call]
        d.add_paragraph("page two content")
        d.save(str(with_break))

        result = parser.parse(with_break)
        assert result.pages_are_synthetic is False
        assert len(result.pages) == 2, result.pages
        assert "page one content" in result.pages[0].text
        assert "page two content" in result.pages[1].text

        # No page break → synthetic pagination.
        no_break = root / "no_break.docx"
        d2 = docx.Document()
        d2.add_paragraph("just one short paragraph")
        d2.save(str(no_break))

        result2 = parser.parse(no_break)
        assert result2.pages_are_synthetic is True
        assert len(result2.pages) == 1

        # Corrupt: garbage bytes with a .docx name.
        corrupt = root / "corrupt.docx"
        corrupt.write_bytes(b"not a docx at all")
        try:
            parser.parse(corrupt)
            raise AssertionError("expected CORRUPT")
        except ParseError as e:
            assert e.kind is ParseErrorKind.CORRUPT

    print("ok: explicit page break, synthetic fallback, and corrupt DOCX all handled correctly")


if __name__ == "__main__":
    _demo()
